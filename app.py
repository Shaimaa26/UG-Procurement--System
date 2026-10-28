import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import pandas as pd
import streamlit as st
from supabase import create_client
from pathlib import Path

LOGO_PATH = Path(__file__).parent / "logo.png"
# =========================================================
# PAGE CONFIG
# =========================================================

# st.set_page_config(
#     page_title="United Group Procurement",
#     page_icon="📊",
#     layout="wide",
# )

# st.title("United Group | Procurement System")

st.set_page_config(
    page_title="United Group Procurement",
    page_icon=str(LOGO_PATH) if LOGO_PATH.exists() else "📊",
    layout="wide",
)

if LOGO_PATH.exists():
    st.logo(str(LOGO_PATH), size="large")

header_logo, header_title = st.columns([1, 6])
with header_logo:
    if LOGO_PATH.exists():
        st.image(str(LOGO_PATH), width=110)
with header_title:
    st.title("United Group | Procurement System")

# =========================================================
# SUPABASE
# =========================================================

@st.cache_resource
def get_supabase():
    url = st.secrets["SUPABASE_URL"]
    key = st.secrets["SUPABASE_KEY"]
    return create_client(url, key)


try:
    supabase = get_supabase()
except Exception as exc:
    st.error(f"Supabase connection failed: {exc}")
    st.stop()

BUCKET = "procurement-documents"


# =========================================================
# TABLE DEFINITIONS
# =========================================================

TABLES = {
    "Suppliers": {
        "table": "suppliers",
        "columns": [
            "id", "supplier_name", "category", "products", "country",
            "contact_person", "contact_title", "email", "phone",
            "payment_terms", "notes",
        ],
        "required": ["supplier_name"],
        "key": "supplier_name",
        "aliases": {
            "supplier": "supplier_name",
            "supplier name": "supplier_name",
            "vendor": "supplier_name",
            "vendor name": "supplier_name",
            "company": "supplier_name",
            "company name": "supplier_name",
            "contact": "contact_person",
            "contact name": "contact_person",
            "telephone": "phone",
            "payment term": "payment_terms",
            "remarks": "notes",
            "comments": "notes",
            "supplier category": "category",
            "products supplied": "products",
            "product list": "products",
            "title": "contact_title",
            "job title": "contact_title",
            "contact title": "contact_title",
            "position": "contact_title",
        },
    },
    "Projects": {
        "table": "projects",
        "columns": [
            "id", "project_name", "customer", "location", "status", "notes",
        ],
        "required": ["project_name"],
        "key": "project_name",
        "aliases": {
            "project": "project_name",
            "project name": "project_name",
            "project title": "project_name",
            "client": "customer",
            "customer name": "customer",
            "site": "location",
            "project status": "status",
            "remarks": "notes",
            "comments": "notes",
        },
    },
    "Products": {
        "table": "products",
        "columns": [
            "id", "product_name", "category", "manufacturer", "model",
            "supplier_id", "unit", "datasheet_url", "notes",
        ],
        "required": ["product_name"],
        "key": "id",
        "aliases": {
            "product": "product_name",
            "product name": "product_name",
            "item": "product_name",
            "item name": "product_name",
            "description": "product_name",
            "product category": "category",
            "type": "category",
            "brand": "manufacturer",
            "make": "manufacturer",
            "part number": "model",
            "model number": "model",
            "supplier": "supplier_id",
            "supplier id": "supplier_id",
            "vendor id": "supplier_id",
            "uom": "unit",
            "unit of measure": "unit",
            "datasheet": "datasheet_url",
            "datasheet link": "datasheet_url",
            "remarks": "notes",
            "comments": "notes",
        },
    },
    
    "Product Accessories": {
        "table": "product_accessories",
        "columns": [
            "id", "product_name", "product_id", "accessory_name",
            "accessory_type", "manufacturer", "model", "quantity", "unit",
            "supplier_id", "unit_price", "currency", "notes",
        ],
        "required": ["accessory_name"],
        "key": "id",
        "aliases": {
            "product": "product_name",
            "product name": "product_name",
            "main product": "product_name",
            "accessory": "accessory_name",
            "accessory name": "accessory_name",
            "item": "accessory_name",
            "description": "accessory_name",
            "type": "accessory_type",
            "category": "accessory_type",
            "accessory category": "accessory_type",
            "brand": "manufacturer",
            "make": "manufacturer",
            "part number": "model",
            "qty": "quantity",
            "uom": "unit",
            "supplier": "supplier_id",
            "price": "unit_price",
            "unit cost": "unit_price",
            "remarks": "notes",
            "comments": "notes",
        },
    },

    #add product prices:
    
        "Product Prices": {
        "table": "product_prices",
        "columns": [
            "id", "product_name", "product_id", "supplier_name",
            "supplier_id", "unit_price", "currency", "price_date",
            "quote_reference", "notes",
        ],
        "required": ["product_id", "unit_price"],
        "key": "id",
        "aliases": {
            "product": "product_name",
            "product name": "product_name",
            "item": "product_name",
            "supplier": "supplier_name",
            "supplier name": "supplier_name",
            "vendor": "supplier_name",
            "price": "unit_price",
            "unit cost": "unit_price",
            "cost": "unit_price",
            "date": "price_date",
            "quote date": "price_date",
            "quote": "quote_reference",
            "quotation": "quote_reference",
            "quote no": "quote_reference",
            "remarks": "notes",
            "comments": "notes",
        },
    },
 #add product prices.
    
    "Inventory": {
        "table": "inventory",
        "columns": [
            "id", "product_id", "serial_number", "quantity",
            "stock_classification", "project_id", "supplier_id",
            "po_number", "receipt_date", "location", "condition", "notes",
        ],
        "required": ["product_id", "quantity", "stock_classification"],
        "key": "id",
        "aliases": {
            "product": "product_id",
            "product id": "product_id",
            "product name": "product_id",
            "serial": "serial_number",
            "serial no": "serial_number",
            "serial number": "serial_number",
            "qty": "quantity",
            "stock quantity": "quantity",
            "classification": "stock_classification",
            "stock type": "stock_classification",
            "project": "project_id",
            "project id": "project_id",
            "supplier": "supplier_id",
            "supplier id": "supplier_id",
            "purchase order": "po_number",
            "po": "po_number",
            "po no": "po_number",
            "date received": "receipt_date",
            "received date": "receipt_date",
            "site": "location",
            "remarks": "notes",
            "comments": "notes",
        },
    },
    "Purchase Orders": {
        "table": "purchase_orders",
        "columns": [
            "id", "po_number", "project_id", "supplier_id", "order_date",
            "currency", "status", "payment_status", "shipping_status", "notes",
        ],
        "required": ["po_number"],
        "key": "po_number",
        "aliases": {
            "po": "po_number",
            "po no": "po_number",
            "po number": "po_number",
            "purchase order": "po_number",
            "purchase order number": "po_number",
            "project": "project_id",
            "project id": "project_id",
            "supplier": "supplier_id",
            "supplier id": "supplier_id",
            "vendor": "supplier_id",
            "date": "order_date",
            "order date": "order_date",
            "payment": "payment_status",
            "shipment status": "shipping_status",
            "remarks": "notes",
            "comments": "notes",
        },
    },
    "Purchase Order Items": {
        "table": "purchase_order_items",
        "columns": [
            "id", "purchase_order_id", "product_id", "quantity",
            "unit_price", "notes",
        ],
        "required": ["purchase_order_id", "product_id", "quantity"],
        "key": "id",
        "aliases": {
            "po id": "purchase_order_id",
            "purchase order id": "purchase_order_id",
            "product": "product_id",
            "product id": "product_id",
            "qty": "quantity",
            "unit cost": "unit_price",
            "price": "unit_price",
            "remarks": "notes",
            "comments": "notes",
        },
    },
}

STOCK_TYPES = ["Project Allocation", "UG Stock", "Project Spare"]


# =========================================================
# COMMON FUNCTIONS
# =========================================================

def fetch_rows(table, limit=5000):
    response = (
        supabase.table(table)
        .select("*")
        .limit(limit)
        .execute()
    )
    return response.data or []


# def clean_value(value):
#     if value is None:
#         return None
#     try:
#         if pd.isna(value):
#             return None
#     except (TypeError, ValueError):
#         pass

#     if isinstance(value, (pd.Timestamp, datetime, date)):
#         return value.isoformat()

#     if isinstance(value, str):
#         value = value.strip()
#         return value if value else None

#     if hasattr(value, "item"):
#         try:
#             return value.item()
#         except Exception:
#             pass

#     return value

def clean_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    if hasattr(value, "item"):
        try:
            value = value.item()
        except Exception:
            pass

    # Whole numbers read as decimals from Excel (4.0) become integers (4)
    if isinstance(value, float) and value.is_integer():
        return int(value)

    return value

def clean_record(record):
    result = {}
    for key, value in record.items():
        value = clean_value(value)
        if value is not None:
            result[key] = value
    return result


def normalize_column(value):
    value = str(value).strip().lower().replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def suggest_mapping(source, config):
    normalized = normalize_column(source)

    for target in config["columns"]:
        if normalized == normalize_column(target):
            return target

    aliases = {
        normalize_column(alias): target
        for alias, target in config["aliases"].items()
    }
    return aliases.get(normalized, "-- Ignore --")


def display_table(table):
    try:
        rows = fetch_rows(table)
        if not rows:
            st.info("No records found.")
            return

        df = pd.DataFrame(rows)
        st.dataframe(df, use_container_width=True, hide_index=True)

        st.download_button(
            "Download CSV",
            df.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"{table}.csv",
            mime="text/csv",
            key=f"csv_{table}",
        )
    except Exception as exc:
        st.error(f"Could not load {table}: {exc}")


def get_choices(table, label_column):
    rows = fetch_rows(table)
    return {
        row["id"]: str(row.get(label_column) or row["id"])
        for row in rows
        if row.get("id") is not None
    }


def add_basic_form(table, fields, title):
    st.subheader(title)

    with st.form(f"form_{table}"):
        values = {}

        for field, kind in fields:
            label = field.replace("_", " ").title()

            if kind == "text":
                values[field] = st.text_input(label)
            elif kind == "textarea":
                values[field] = st.text_area(label)
            elif kind == "number":
                values[field] = st.number_input(
                    label, min_value=0.0, value=0.0
                )
            elif kind == "date":
                values[field] = st.date_input(label)
            elif kind == "supplier":
                choices = get_choices("suppliers", "supplier_name")
                options = [None] + list(choices.keys())
                values[field] = st.selectbox(
                    label, options,
                    format_func=lambda x, c=choices: "-- None --" if x is None else c[x],
                )
            elif kind == "project":
                choices = get_choices("projects", "project_name")
                options = [None] + list(choices.keys())
                values[field] = st.selectbox(
                    label, options,
                    format_func=lambda x, c=choices: "-- None --" if x is None else c[x],
                )
            elif kind == "product":
                choices = get_choices("products", "product_name")
                options = [None] + list(choices.keys())
                values[field] = st.selectbox(
                    label, options,
                    format_func=lambda x, c=choices: "-- None --" if x is None else c[x],
                )
            elif kind == "po":
                choices = get_choices("purchase_orders", "po_number")
                options = [None] + list(choices.keys())
                values[field] = st.selectbox(
                    label, options,
                    format_func=lambda x, c=choices: "-- None --" if x is None else c[x],
                )
            elif kind == "stock":
                values[field] = st.selectbox(label, STOCK_TYPES)

        submitted = st.form_submit_button("Save record")

    if submitted:
        payload = {}
        for field, kind in fields:
            value = values[field]
            if isinstance(value, (date, datetime)):
                value = value.isoformat()
            value = clean_value(value)
            if value is not None:
                payload[field] = value

        try:
            supabase.table(table).insert(payload).execute()
            st.success("Record saved.")
            st.rerun()
        except Exception as exc:
            st.error(f"Could not save record: {exc}")



def build_template(label, include_id=False):
    """Create an Excel template for one section, with instructions
    and reference sheets that list the IDs needed for linked fields."""
    config = TABLES[label]
    columns = [
        c for c in config["columns"] if include_id or c != "id"
    ]

    notes = {
        "id": "Only needed when updating existing records.",
        "supplier_id": "Use the supplier ID (see 'Suppliers ref' sheet).",
        "project_id": "Use the project ID (see 'Projects ref' sheet).",
        "product_id": "Use the product ID (see 'Products ref' sheet).",
        "purchase_order_id": "Use the PO ID (see 'PO ref' sheet).",
        "stock_classification": "Use exactly: " + ", ".join(STOCK_TYPES),
        "category": "Choose from the dropdown (see the 'categories ref' tabs). Same spelling avoids duplicates.",
        "receipt_date": "Date format: YYYY-MM-DD",
        "order_date": "Date format: YYYY-MM-DD",
        "quantity": "Number, zero or more.",
        "unit_price": "Number, zero or more.",
        "product_name": "Type the exact product name (or use product_id from the 'Products ref' sheet).",
        "accessory_type": "For example: Mounting bracket, Cable, Enclosure.",
        "supplier_name": "Type the exact supplier name saved in Suppliers.",
        "currency": "For example: USD, EUR, EGP. Use the same spelling every time.",
        "price_date": "Date of the quote or price list: YYYY-MM-DD",
    }

    instructions = pd.DataFrame({
        "Column": columns,
        "Required": [
            "YES" if c in config["required"] else "no" for c in columns
        ],
        "Notes": [notes.get(c, "") for c in columns],
    })

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(columns=columns).to_excel(
            writer, sheet_name="Template", index=False
        )
        instructions.to_excel(writer, sheet_name="Instructions", index=False)

        references = [
            ("Suppliers ref", "suppliers", ["id", "supplier_name"]),
            ("Projects ref", "projects", ["id", "project_name"]),
            ("Products ref", "products", ["id", "product_name", "model"]),
            ("PO ref", "purchase_orders", ["id", "po_number"]),
        ]

        for sheet_name, ref_table, ref_columns in references:
            try:
                rows = fetch_rows(ref_table)
                ref_df = pd.DataFrame(rows)
                for col in ref_columns:
                    if col not in ref_df.columns:
                        ref_df[col] = None
                ref_df = ref_df[ref_columns]
            except Exception:
                ref_df = pd.DataFrame(columns=ref_columns)
            ref_df.to_excel(writer, sheet_name=sheet_name, index=False)



        #Adding category reference
                # Categories already used in the database (to avoid duplicates)
        from openpyxl.worksheet.datavalidation import DataValidation

        category_sheets = [
            ("Supplier categories ref", "suppliers"),
            ("Product categories ref", "products"),
        ]
        category_ranges = {}

        for cat_sheet, source_table in category_sheets:
            try:
                source = pd.DataFrame(fetch_rows(source_table))
                if "category" in source.columns:
                    cats = source["category"].dropna().astype(str).str.strip()
                    cats = cats[cats != ""]
                    cat_df = (
                        cats.value_counts()
                        .rename_axis("category")
                        .reset_index(name="records")
                        .sort_values(
                            "category", key=lambda s: s.str.lower()
                        )
                    )
                else:
                    cat_df = pd.DataFrame(columns=["category", "records"])
            except Exception:
                cat_df = pd.DataFrame(columns=["category", "records"])

            cat_df.to_excel(writer, sheet_name=cat_sheet, index=False)
            category_ranges[source_table] = (cat_sheet, len(cat_df))

        # Dropdown on the 'category' column of the Template sheet
        if "category" in columns and config["table"] in category_ranges:
            ref_sheet, cat_count = category_ranges[config["table"]]
            if cat_count > 0:
                template_ws = writer.book["Template"]
                cat_letter = template_ws.cell(
                    row=1, column=columns.index("category") + 1
                ).column_letter
                dropdown = DataValidation(
                    type="list",
                    formula1=f"='{ref_sheet}'!$A$2:$A${cat_count + 1}",
                    allow_blank=True,
                    showErrorMessage=True,
                    errorStyle="warning",
                    errorTitle="New category",
                    error=(
                        "This category is not in the list. Use an existing "
                        "category unless you really need a new one."
                    ),
                )
                template_ws.add_data_validation(dropdown)
                dropdown.add(f"{cat_letter}2:{cat_letter}1000")

        
        # Light formatting: bold headers, required columns in yellow
        from openpyxl.styles import Font, PatternFill

        bold = Font(bold=True)
        yellow = PatternFill("solid", fgColor="FFF2CC")

        for ws in writer.book.worksheets:
            for cell in ws[1]:
                cell.font = bold
                if (
                    ws.title == "Template"
                    and cell.value in config["required"]
                ):
                    cell.fill = yellow
            for column_cells in ws.columns:
                width = max(
                    len(str(c.value)) if c.value is not None else 0
                    for c in column_cells
                )
                ws.column_dimensions[column_cells[0].column_letter].width = (
                    min(max(width + 2, 14), 45)
                )

    return output.getvalue()


# =========================================================
# LOGIN
# =========================================================

if not st.session_state.get("authenticated", False):
    st.subheader("Login")

    with st.form("login"):
        password = st.text_input("Application password", type="password")
        submitted = st.form_submit_button("Login")

    if submitted:
        if password == st.secrets.get("APP_PASSWORD", ""):
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("Incorrect password.")

    st.stop()


# =========================================================
# NAVIGATION
# =========================================================

with st.sidebar:
    page = st.radio(
        "Navigation",
        [
            "Dashboard",
            "Search",
            "Suppliers",
            "Projects",
            "Products",
            "Accessories",
            "Price List",
            "Inventory",
            "Purchase Orders",
            "Purchase Order Items",
            "Excel Import",
            "Documents",
        ],
    )

    if st.button("Logout"):
        st.session_state["authenticated"] = False
        st.rerun()


# =========================================================
# DASHBOARD
# =========================================================

if page == "Dashboard":
    st.subheader("Overview")

    dashboard_sections = [
        ("Suppliers", "suppliers"),
        ("Projects", "projects"),
        ("Products", "products"),
        ("Inventory", "inventory"),
        ("Purchase Orders", "purchase_orders"),
        ("Purchase Order Items", "purchase_order_items"),
    ]

    if "dashboard_section" not in st.session_state:
        st.session_state["dashboard_section"] = "purchase_orders"

    def count_rows(table_name):
        try:
            result = (
                supabase.table(table_name)
                .select("id", count="exact")
                .limit(1)
                .execute()
            )
            if result.count is not None:
                return result.count
            return len(fetch_rows(table_name))
        except Exception:
            return "Error"

    # Counters with a button under each one
    for start in range(0, len(dashboard_sections), 4):
        row_cols = st.columns(4)
        for col, (label, table_name) in zip(
            row_cols, dashboard_sections[start:start + 4]
        ):
            is_selected = st.session_state["dashboard_section"] == table_name
            col.metric(label, count_rows(table_name))
            if col.button(
                "View records",
                key=f"dash_btn_{table_name}",
                type="primary" if is_selected else "secondary",
                use_container_width=True,
            ):
                st.session_state["dashboard_section"] = table_name
                st.rerun()

    # Records of the selected section, shown on the same page
    selected_table = st.session_state["dashboard_section"]
    selected_label = next(
        (lbl for lbl, tbl in dashboard_sections if tbl == selected_table),
        selected_table,
    )

    st.divider()
    st.subheader(selected_label)

    try:
        dash_rows = fetch_rows(selected_table)

        if not dash_rows:
            st.info("No records found.")
        else:
            dash_df = pd.DataFrame(dash_rows)

            dash_keyword = st.text_input(
                "Search in these records",
                key=f"dash_search_{selected_table}",
            )

            if dash_keyword:
                dash_mask = (
                    dash_df.astype(str)
                    .apply(
                        lambda c: c.str.contains(
                            dash_keyword, case=False, na=False, regex=False
                        )
                    )
                    .any(axis=1)
                )
                dash_df = dash_df[dash_mask]

            st.caption(f"{len(dash_df)} record(s)")
            st.dataframe(dash_df, use_container_width=True, hide_index=True)

            st.download_button(
                "Download CSV",
                dash_df.to_csv(index=False).encode("utf-8-sig"),
                file_name=f"{selected_table}.csv",
                mime="text/csv",
                key=f"dash_csv_{selected_table}",
            )
    except Exception as exc:
        st.error(f"Could not load {selected_label}: {exc}")




# #if page == "Dashboard":
#     st.subheader("Overview")
#     metric_cols = st.columns(5)

#     for col, label, table in zip(
#         metric_cols,
#         ["Suppliers", "Projects", "Products", "Inventory", "Purchase Orders"],
#         ["suppliers", "projects", "products", "inventory", "purchase_orders"],
#     ):
#         try:
#             col.metric(label, len(fetch_rows(table)))
#         except Exception:
#             col.metric(label, "Error")

#     st.subheader("Purchase Orders")
#     #display_table("purchase_orders")


# elif page == "Suppliers":
#     add_basic_form(
#         "suppliers",
#         [
#             ("supplier_name", "text"), ("category", "text"),
#             ("products", "textarea"), ("country", "text"),
#             ("contact_person", "text"), ("contact_title", "text"),
#             ("email", "text"), ("phone", "text"),
#             ("payment_terms", "text"), ("notes", "textarea"),
#         ],
#         "Add Supplier",
#     )
#     display_table("suppliers")


elif page == "Projects":
    add_basic_form(
        "projects",
        [
            ("project_name", "text"), ("customer", "text"),
            ("location", "text"), ("status", "text"),
            ("notes", "textarea"),
        ],
        "Add Project",
    )
    display_table("projects")


elif page == "Products":
    add_basic_form(
        "products",
        [
            ("product_name", "text"), ("category", "text"),
            ("manufacturer", "text"), ("model", "text"),
            ("supplier_id", "supplier"), ("unit", "text"),
            ("datasheet_url", "text"), ("notes", "textarea"),
        ],
        "Add Product",
    )
    display_table("products")


#added with product accessories:

elif page == "Accessories":
    st.subheader("Product Accessories")

    acc_products = {
        row["id"]: str(row.get("product_name") or row["id"])
        for row in fetch_rows("products")
        if row.get("id") is not None
    }

    with st.expander("Add a new accessory"):
        add_basic_form(
            "product_accessories",
            [
                ("product_id", "product"), ("accessory_name", "text"),
                ("accessory_type", "text"), ("manufacturer", "text"),
                ("model", "text"), ("quantity", "number"),
                ("unit", "text"), ("supplier_id", "supplier"),
                ("unit_price", "number"), ("currency", "text"), 
                ("notes", "textarea"),
            ],
            "Add Accessory",
        )

    st.divider()

    scope = st.radio(
        "Show accessories for",
        ["One product", "All products"],
        horizontal=True,
    )

    selected_product = None
    if scope == "One product":
        selected_product = st.selectbox(
            "Product",
            [None] + list(acc_products.keys()),
            format_func=lambda x: (
                "-- Select a product --" if x is None else acc_products[x]
            ),
        )
        if selected_product is None:
            st.info("Select a product to see its accessories.")
            st.stop()

    acc_rows = fetch_rows("product_accessories")
    if not acc_rows:
        st.info("No accessories saved yet.")
        st.stop()

    acc_df = pd.DataFrame(acc_rows)

    if selected_product is not None:
        acc_df = acc_df[acc_df["product_id"] == selected_product]

    acc_df.insert(
        1, "product_name", acc_df["product_id"].map(acc_products)
    )

    f1, f2 = st.columns(2)

    type_values = sorted(
        acc_df["accessory_type"].dropna().astype(str).unique()
    ) if "accessory_type" in acc_df.columns else []

    chosen_types = f1.multiselect("Accessory type", type_values)
    acc_keyword = f2.text_input("Search in accessories")

    if chosen_types:
        acc_df = acc_df[acc_df["accessory_type"].astype(str).isin(chosen_types)]

    if acc_keyword:
        acc_mask = (
            acc_df.astype(str)
            .apply(
                lambda c: c.str.contains(
                    acc_keyword, case=False, na=False, regex=False
                )
            )
            .any(axis=1)
        )
        acc_df = acc_df[acc_mask]

    m1, m2 = st.columns(2)
    m1.metric("Accessories", len(acc_df))
    m2.metric("Products covered", acc_df["product_id"].nunique())

    st.dataframe(acc_df, use_container_width=True, hide_index=True)

    st.download_button(
        "Download accessories (CSV)",
        acc_df.to_csv(index=False).encode("utf-8-sig"),
        file_name="product_accessories.csv",
        mime="text/csv",
        key="csv_accessories",
    )


#added with product accessories.

#Add product price:
elif page == "Price List":
    st.subheader("Product Price List")

    price_products = {
        row["id"]: str(row.get("product_name") or row["id"])
        for row in fetch_rows("products")
        if row.get("id") is not None
    }
    price_suppliers = get_choices("suppliers", "supplier_name")

    with st.expander("Add a price"):
        add_basic_form(
            "product_prices",
            [
                ("product_id", "product"), ("supplier_id", "supplier"),
                ("unit_price", "number"), ("currency", "text"),
                ("price_date", "date"), ("quote_reference", "text"),
                ("notes", "textarea"),
            ],
            "Add Price",
        )
# add button to just complete price for exisitng products
    with st.expander("Download price template for existing products"):

        template_scope = st.radio(
            "Products to include",
            ["Only products without a price", "All products"],
            horizontal=True,
            key="price_template_scope",
        )

        tpl_products = pd.DataFrame(fetch_rows("products"))
        tpl_prices = pd.DataFrame(fetch_rows("product_prices"))

        if tpl_products.empty:
            st.info("No products yet.")
        else:
            for col in ["id", "product_name", "category", "manufacturer",
                        "model", "unit"]:
                if col not in tpl_products.columns:
                    tpl_products[col] = None
            if "product_id" not in tpl_prices.columns:
                tpl_prices["product_id"] = None

            def tpl_key(v):
                if pd.isna(v):
                    return None
                if isinstance(v, float) and v.is_integer():
                    return str(int(v))
                return str(v)

            if template_scope == "Only products without a price":
                priced_ids = {
                    tpl_key(v) for v in tpl_prices["product_id"].dropna()
                }
                tpl_products = tpl_products[
                    ~tpl_products["id"].map(tpl_key).isin(priced_ids)
                ]

            if tpl_products.empty:
                st.info("All products already have a price.")
            else:
                tpl_df = pd.DataFrame({
                    "product_id": tpl_products["id"],
                    "product_name": tpl_products["product_name"],
                    "category": tpl_products["category"],
                    "manufacturer": tpl_products["manufacturer"],
                    "model": tpl_products["model"],
                    "unit": tpl_products["unit"],
                    "supplier_name": None,
                    "unit_price": None,
                    "currency": None,
                    "price_date": None,
                    "quote_reference": None,
                    "notes": None,
                }).sort_values("product_name", na_position="last")

                sup_names = sorted(
                    str(n) for n in price_suppliers.values()
                )

                tpl_buffer = io.BytesIO()
                with pd.ExcelWriter(tpl_buffer, engine="openpyxl") as writer:
                    tpl_df.to_excel(writer, sheet_name="Template", index=False)
                    pd.DataFrame({"supplier_name": sup_names}).to_excel(
                        writer, sheet_name="Suppliers ref", index=False
                    )

                    from openpyxl.styles import Font, PatternFill
                    from openpyxl.worksheet.datavalidation import DataValidation

                    ws = writer.book["Template"]
                    cols = list(tpl_df.columns)

                    for cell in ws[1]:
                        cell.font = Font(bold=True)
                        if cell.value in ("supplier_name", "unit_price"):
                            cell.fill = PatternFill("solid", fgColor="FFF2CC")

                    for column_cells in ws.columns:
                        width = max(
                            len(str(c.value)) if c.value is not None else 0
                            for c in column_cells
                        )
                        ws.column_dimensions[
                            column_cells[0].column_letter
                        ].width = min(max(width + 2, 14), 45)

                    if sup_names:
                        letter = ws.cell(
                            row=1, column=cols.index("supplier_name") + 1
                        ).column_letter
                        dropdown = DataValidation(
                            type="list",
                            formula1=f"='Suppliers ref'!$A$2:$A${len(sup_names) + 1}",
                            allow_blank=True,
                            showErrorMessage=True,
                            errorStyle="warning",
                            errorTitle="Unknown supplier",
                            error="This supplier is not saved yet. Add it in Suppliers first.",
                        )
                        ws.add_data_validation(dropdown)
                        dropdown.add(f"{letter}2:{letter}{len(tpl_df) + 1}")

                st.caption(f"{len(tpl_df)} product(s) in the file")
                st.download_button(
                    "Download price template (Excel)",
                    tpl_buffer.getvalue(),
                    file_name="product_prices_to_complete.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="price_template_download",
                )
                st.caption(
                    "Fill supplier_name, unit_price, currency and price_date "
                    "(yellow headers are the main ones). Keep product_id and "
                    "product_name as they are."
                )

    
# add button to add pricing for existing products

# adding accessories with products in 1 file of pricing
    with st.expander("Download accessories price template"):
        at_scope = st.radio(
            "Accessories to include",
            ["Only accessories without a price", "All accessories"],
            horizontal=True,
            key="acc_price_scope",
        )

        at_acc = pd.DataFrame(fetch_rows("product_accessories"))

        def at_key(v):
            if pd.isna(v):
                return None
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)

        at_products = {
            str(row["id"]): str(row.get("product_name") or row["id"])
            for row in fetch_rows("products")
            if row.get("id") is not None
        }

        if at_acc.empty:
            st.info("No accessories yet. Add them on the Accessories page first.")
        else:
            for col in ["id", "product_id", "accessory_name", "accessory_type",
                        "manufacturer", "model", "quantity", "unit",
                        "unit_price"]:
                if col not in at_acc.columns:
                    at_acc[col] = None

            if at_scope == "Only accessories without a price":
                at_acc = at_acc[
                    pd.to_numeric(at_acc["unit_price"], errors="coerce").isna()
                ]

            if at_acc.empty:
                st.info("All accessories already have a price.")
            else:
                at_df = pd.DataFrame({
                    "id": at_acc["id"],
                    "device": at_acc["product_id"].map(
                        lambda v: at_products.get(at_key(v))
                    ),
                    "accessory_type": at_acc["accessory_type"],
                    "accessory_name": at_acc["accessory_name"],
                    "manufacturer": at_acc["manufacturer"],
                    "model": at_acc["model"],
                    "quantity": at_acc["quantity"],
                    "unit": at_acc["unit"],
                    "unit_price": None,
                }).sort_values(
                    ["device", "accessory_type"], na_position="last"
                )

                at_buffer = io.BytesIO()
                with pd.ExcelWriter(at_buffer, engine="openpyxl") as writer:
                    at_df.to_excel(writer, sheet_name="Template", index=False)

                    from openpyxl.styles import Font, PatternFill

                    ws = writer.book["Template"]
                    for cell in ws[1]:
                        cell.font = Font(bold=True)
                        if cell.value == "unit_price":
                            cell.fill = PatternFill("solid", fgColor="FFF2CC")
                    for column_cells in ws.columns:
                        width = max(
                            len(str(c.value)) if c.value is not None else 0
                            for c in column_cells
                        )
                        ws.column_dimensions[
                            column_cells[0].column_letter
                        ].width = min(max(width + 2, 12), 45)

                st.caption(f"{len(at_df)} accessory line(s) in the file")
                st.download_button(
                    "Download accessories price template (Excel)",
                    at_buffer.getvalue(),
                    file_name="accessories_prices_to_complete.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="acc_price_download",
                )
                st.caption(
                    "Fill unit_price only. Keep the id column. To import: "
                    "Excel Import, destination Product Accessories, mode "
                    "Update existing, match by id."
                )


###################################################################
    
    with st.expander("Download complete devices (product + accessories + prices)"):
        cd_products = pd.DataFrame(fetch_rows("products"))
        cd_acc = pd.DataFrame(fetch_rows("product_accessories"))
        cd_prices = pd.DataFrame(fetch_rows("product_prices"))

        def cd_key(v):
            if pd.isna(v):
                return None
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)

        cd_suppliers = {
            str(k): v
            for k, v in get_choices("suppliers", "supplier_name").items()
        }

        if cd_products.empty:
            st.info("No products yet.")
        else:
            for col in ["id", "product_name", "category", "manufacturer",
                        "model", "unit"]:
                if col not in cd_products.columns:
                    cd_products[col] = None
            for col in ["product_id", "accessory_name", "accessory_type",
                        "manufacturer", "model", "quantity", "unit",
                        "unit_price", "supplier_id"]:
                if col not in cd_acc.columns:
                    cd_acc[col] = None
            for col in ["product_id", "supplier_id", "unit_price", "currency",
                        "price_date", "created_at"]:
                if col not in cd_prices.columns:
                    cd_prices[col] = None

            o1, o2 = st.columns(2)
            price_basis = o1.radio(
                "Device price",
                ["Latest price", "Lowest price"],
                horizontal=True,
                key="cd_basis",
            )
            only_with_acc = o2.checkbox(
                "Only products that have accessories",
                value=True,
                key="cd_only",
            )
            cd_selected = st.multiselect(
                "Products (leave empty for all)",
                sorted(cd_products["product_name"].dropna().astype(str).unique()),
                key="cd_selected",
            )

            cd_products["_k"] = cd_products["id"].map(cd_key)
            cd_acc["_k"] = cd_acc["product_id"].map(cd_key)
            cd_prices["_k"] = cd_prices["product_id"].map(cd_key)

            cd_prices["unit_price"] = pd.to_numeric(
                cd_prices["unit_price"], errors="coerce"
            )
            cd_prices = cd_prices.dropna(subset=["unit_price"]).copy()
            cd_prices["supplier_name"] = cd_prices["supplier_id"].map(
                lambda v: cd_suppliers.get(cd_key(v))
            )

            if cd_prices.empty:
                best = pd.DataFrame(columns=[
                    "_k", "device_price", "currency", "device_supplier",
                ])
            else:
                if price_basis == "Latest price":
                    best = (
                        cd_prices.sort_values(
                            ["price_date", "created_at"], na_position="first"
                        ).drop_duplicates("_k", keep="last")
                    )
                else:
                    best = (
                        cd_prices.sort_values("unit_price")
                        .drop_duplicates("_k", keep="first")
                    )
                best = best[
                    ["_k", "unit_price", "currency", "supplier_name"]
                ].rename(columns={
                    "unit_price": "device_price",
                    "supplier_name": "device_supplier",
                })

            devices = cd_products.merge(best, on="_k", how="left")

            if cd_selected:
                devices = devices[
                    devices["product_name"].astype(str).isin(cd_selected)
                ]
            if only_with_acc:
                devices = devices[devices["_k"].isin(set(cd_acc["_k"].dropna()))]

            if devices.empty:
                st.info("No products match these options.")
            else:
                detail_rows = []
                summary_rows = []

                for _, dev in devices.sort_values(
                    "product_name", na_position="last"
                ).iterrows():
                    dev_price = dev.get("device_price")
                    dev_currency = dev.get("currency")
                    dev_has_price = pd.notna(dev_price)

                    detail_rows.append({
                        "device": dev["product_name"],
                        "item_type": "Device",
                        "item_name": dev["product_name"],
                        "accessory_type": None,
                        "manufacturer": dev["manufacturer"],
                        "model": dev["model"],
                        "quantity": 1,
                        "unit": dev["unit"],
                        "unit_price": dev_price if dev_has_price else None,
                        "currency": dev_currency,
                        "line_total": dev_price if dev_has_price else None,
                        "price_supplier": dev.get("device_supplier"),
                    })

                    acc_total = 0.0
                    acc_count = 0
                    unpriced = 0 if dev_has_price else 1

                    for _, acc in cd_acc[cd_acc["_k"] == dev["_k"]].iterrows():
                        qty = pd.to_numeric(acc["quantity"], errors="coerce")
                        price = pd.to_numeric(acc["unit_price"], errors="coerce")
                        qty_for_total = 1 if pd.isna(qty) else qty

                        if pd.notna(price):
                            line_total = qty_for_total * price
                            acc_total += line_total
                        else:
                            line_total = None
                            unpriced += 1

                        acc_count += 1
                        detail_rows.append({
                            "device": dev["product_name"],
                            "item_type": "Accessory",
                            "item_name": acc["accessory_name"],
                            "accessory_type": acc["accessory_type"],
                            "manufacturer": acc["manufacturer"],
                            "model": acc["model"],
                            "quantity": None if pd.isna(qty) else qty,
                            "unit": acc["unit"],
                            "unit_price": None if pd.isna(price) else price,
                            "currency": dev_currency,
                            "line_total": line_total,
                            "price_supplier": cd_suppliers.get(
                                cd_key(acc["supplier_id"])
                            ),
                        })

                    summary_rows.append({
                        "device": dev["product_name"],
                        "category": dev["category"],
                        "manufacturer": dev["manufacturer"],
                        "model": dev["model"],
                        "price_supplier": dev.get("device_supplier"),
                        "device_price": dev_price if dev_has_price else None,
                        "accessories_count": acc_count,
                        "accessories_total": acc_total if acc_count else None,
                        "complete_total": (
                            (dev_price if dev_has_price else 0) + acc_total
                        ),
                        "currency": dev_currency,
                        "items_without_price": unpriced,
                    })

                detail_df = pd.DataFrame(detail_rows)
                summary_df = pd.DataFrame(summary_rows)

                cd_buffer = io.BytesIO()
                with pd.ExcelWriter(cd_buffer, engine="openpyxl") as writer:
                    detail_df.to_excel(
                        writer, sheet_name="Complete devices", index=False
                    )
                    summary_df.to_excel(
                        writer, sheet_name="Summary", index=False
                    )

                    from openpyxl.styles import Font, PatternFill

                    device_fill = PatternFill("solid", fgColor="E7F0FA")

                    for ws in writer.book.worksheets:
                        for cell in ws[1]:
                            cell.font = Font(bold=True)
                        for column_cells in ws.columns:
                            width = max(
                                len(str(c.value)) if c.value is not None else 0
                                for c in column_cells
                            )
                            ws.column_dimensions[
                                column_cells[0].column_letter
                            ].width = min(max(width + 2, 12), 45)

                    ws_detail = writer.book["Complete devices"]
                    for row in ws_detail.iter_rows(min_row=2):
                        if row[1].value == "Device":
                            for cell in row:
                                cell.fill = device_fill
                                cell.font = Font(bold=True)

                st.caption(
                    f"{len(summary_df)} device(s), "
                    f"{len(detail_df) - len(summary_df)} accessory line(s)"
                )
                st.dataframe(summary_df, use_container_width=True, hide_index=True)

                st.download_button(
                    "Download complete devices (Excel)",
                    cd_buffer.getvalue(),
                    file_name="complete_devices_with_accessories.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="cd_download",
                )
        

# adding accessories with products in 1 file of pricing.       
        
# updated to import one time in one file for products and accessories

    with st.expander("Download combined price file (devices + accessories)"):
        cp_scope = st.radio(
            "Items to include",
            ["Only items without a price", "All items"],
            horizontal=True,
            key="cp_scope",
        )
        cp_only_acc = st.checkbox(
            "Only products that have accessories",
            value=False,
            key="cp_only_acc",
        )

        cp_products = pd.DataFrame(fetch_rows("products"))
        cp_acc = pd.DataFrame(fetch_rows("product_accessories"))
        cp_prices = pd.DataFrame(fetch_rows("product_prices"))
        cp_sup = {
            str(k): v
            for k, v in get_choices("suppliers", "supplier_name").items()
        }

        def cp_key(v):
            if v is None:
                return None
            try:
                if pd.isna(v):
                    return None
            except (TypeError, ValueError):
                pass
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)

        if cp_products.empty:
            st.info("No products yet.")
        else:
            for col in ["id", "product_name", "manufacturer", "model", "unit"]:
                if col not in cp_products.columns:
                    cp_products[col] = None
            for col in ["id", "product_id", "accessory_name", "accessory_type",
                        "manufacturer", "model", "quantity", "unit",
                        "unit_price", "supplier_id"]:
                if col not in cp_acc.columns:
                    cp_acc[col] = None
            if "product_id" not in cp_prices.columns:
                cp_prices["product_id"] = None

            cp_acc["_k"] = cp_acc["product_id"].map(cp_key)
            priced_products = {
                cp_key(v) for v in cp_prices["product_id"].dropna()
            }

            cp_rows = []

            for _, prod in cp_products.sort_values(
                "product_name", na_position="last"
            ).iterrows():
                pk = cp_key(prod["id"])
                prod_acc = cp_acc[cp_acc["_k"] == pk]

                if cp_only_acc and prod_acc.empty:
                    continue

                if cp_scope == "Only items without a price":
                    include_device = pk not in priced_products
                    prod_acc = prod_acc[
                        pd.to_numeric(prod_acc["unit_price"], errors="coerce").isna()
                    ]
                else:
                    include_device = True

                if not include_device and prod_acc.empty:
                    continue

                if include_device:
                    cp_rows.append({
                        "item_type": "Device",
                        "device": prod["product_name"],
                        "product_id": prod["id"],
                        "accessory_id": None,
                        "accessory_type": None,
                        "item_name": prod["product_name"],
                        "manufacturer": prod["manufacturer"],
                        "model": prod["model"],
                        "quantity": None,
                        "unit": prod["unit"],
                        "supplier_name": None,
                        "unit_price": None,
                        "currency": None,
                        "price_date": None,
                        "quote_reference": None,
                        "notes": None,
                    })

                for _, acc in prod_acc.iterrows():
                    cp_rows.append({
                        "item_type": "Accessory",
                        "device": prod["product_name"],
                        "product_id": prod["id"],
                        "accessory_id": acc["id"],
                        "accessory_type": acc["accessory_type"],
                        "item_name": acc["accessory_name"],
                        "manufacturer": acc["manufacturer"],
                        "model": acc["model"],
                        "quantity": acc["quantity"],
                        "unit": acc["unit"],
                        "supplier_name": cp_sup.get(cp_key(acc["supplier_id"])),
                        "unit_price": None,
                        "currency": None,
                        "price_date": None,
                        "quote_reference": None,
                        "notes": None,
                    })

            if not cp_rows:
                st.info("Nothing to include. All items already have a price.")
            else:
                cp_df_out = pd.DataFrame(cp_rows)
                sup_list = sorted(str(n) for n in cp_sup.values())

                cp_buffer = io.BytesIO()
                with pd.ExcelWriter(cp_buffer, engine="openpyxl") as writer:
                    cp_df_out.to_excel(writer, sheet_name="Template", index=False)
                    pd.DataFrame({"supplier_name": sup_list}).to_excel(
                        writer, sheet_name="Suppliers ref", index=False
                    )

                    from openpyxl.styles import Font, PatternFill
                    from openpyxl.worksheet.datavalidation import DataValidation

                    ws = writer.book["Template"]
                    out_cols = list(cp_df_out.columns)
                    yellow_cols = ("supplier_name", "unit_price",
                                   "currency", "price_date")

                    for cell in ws[1]:
                        cell.font = Font(bold=True)
                        if cell.value in yellow_cols:
                            cell.fill = PatternFill("solid", fgColor="FFF2CC")

                    for row in ws.iter_rows(min_row=2):
                        if row[0].value == "Device":
                            for cell in row:
                                cell.fill = PatternFill("solid", fgColor="E7F0FA")
                                cell.font = Font(bold=True)

                    for column_cells in ws.columns:
                        width = max(
                            len(str(c.value)) if c.value is not None else 0
                            for c in column_cells
                        )
                        ws.column_dimensions[
                            column_cells[0].column_letter
                        ].width = min(max(width + 2, 12), 45)

                    if sup_list:
                        letter = ws.cell(
                            row=1, column=out_cols.index("supplier_name") + 1
                        ).column_letter
                        dropdown = DataValidation(
                            type="list",
                            formula1=f"='Suppliers ref'!$A$2:$A${len(sup_list) + 1}",
                            allow_blank=True,
                            showErrorMessage=True,
                            errorStyle="warning",
                            errorTitle="Unknown supplier",
                            error="This supplier is not saved yet. Add it in Suppliers first.",
                        )
                        ws.add_data_validation(dropdown)
                        dropdown.add(f"{letter}2:{letter}{len(cp_df_out) + 1}")

                st.caption(f"{len(cp_df_out)} row(s) in the file")
                st.download_button(
                    "Download combined price file (Excel)",
                    cp_buffer.getvalue(),
                    file_name="devices_and_accessories_prices.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                    key="cp_download",
                )
                st.caption(
                    "Fill supplier_name, unit_price, currency and price_date. "
                    "Leave rows without a price empty. Do not change "
                    "item_type, product_id or accessory_id."
                )


 #import combined file once::
        
    with st.expander("Import combined price file (devices + accessories)"):
        cp_file = st.file_uploader(
            "Upload the completed combined price file",
            type=["xlsx", "xls"],
            key="cp_upload",
        )

        if cp_file:
            try:
                def imp_key(v):
                    if v is None:
                        return None
                    try:
                        if pd.isna(v):
                            return None
                    except (TypeError, ValueError):
                        pass
                    if isinstance(v, float) and v.is_integer():
                        return str(int(v))
                    return str(v)

                imp_bytes = cp_file.getvalue()
                imp_excel = pd.ExcelFile(io.BytesIO(imp_bytes))
                imp_sheet = (
                    "Template" if "Template" in imp_excel.sheet_names
                    else imp_excel.sheet_names[0]
                )
                imp_df = pd.read_excel(
                    io.BytesIO(imp_bytes), sheet_name=imp_sheet, dtype=object
                )
                imp_df.columns = [
                    normalize_column(c).replace(" ", "_") for c in imp_df.columns
                ]

                missing_cols = [
                    c for c in ("item_type", "unit_price")
                    if c not in imp_df.columns
                ]

                if missing_cols:
                    st.error(
                        "This file is missing the column(s): "
                        + ", ".join(missing_cols)
                        + ". Use the combined price file downloaded above."
                    )
                else:
                    imp_suppliers = {
                        str(n).strip().lower(): k
                        for k, n in get_choices(
                            "suppliers", "supplier_name"
                        ).items()
                    }
                    imp_product_names = {}
                    imp_product_ids = set()
                    for k, n in get_choices("products", "product_name").items():
                        imp_product_names.setdefault(
                            str(n).strip().lower(), []
                        ).append(k)
                        imp_product_ids.add(imp_key(k))

                    imp_acc_ids = {
                        imp_key(r.get("id"))
                        for r in fetch_rows("product_accessories")
                    }

                    existing_prices = set()
                    for r in fetch_rows("product_prices"):
                        try:
                            p = round(float(r.get("unit_price")), 6)
                        except (TypeError, ValueError):
                            p = None
                        existing_prices.add((
                            imp_key(r.get("product_id")),
                            imp_key(r.get("supplier_id")),
                            p,
                            (str(r.get("price_date") or "")[:10]) or None,
                        ))

                    device_records = []
                    accessory_updates = []
                    imp_errors = []
                    skipped_empty = 0
                    skipped_duplicates = 0
                    seen_prices = set()

                    for idx, row in imp_df.iterrows():
                        row_no = idx + 2
                        price_raw = clean_value(row.get("unit_price"))

                        if price_raw is None:
                            skipped_empty += 1
                            continue

                        item_type = str(
                            clean_value(row.get("item_type")) or ""
                        ).strip().lower()

                        try:
                            price = float(Decimal(str(price_raw)))
                            if price < 0:
                                raise ValueError("negative")
                        except Exception:
                            imp_errors.append({
                                "Excel row": row_no,
                                "Errors": "unit_price must be a number, zero or more",
                            })
                            continue

                        supplier_text = clean_value(row.get("supplier_name"))
                        supplier_id = None
                        if supplier_text is not None:
                            supplier_id = imp_suppliers.get(
                                str(supplier_text).strip().lower()
                            )
                            if supplier_id is None:
                                imp_errors.append({
                                    "Excel row": row_no,
                                    "Errors": f"Supplier not found: {supplier_text}",
                                })
                                continue

                        if item_type == "device":
                            product_id = clean_value(row.get("product_id"))

                            if product_id is None:
                                device_name = clean_value(row.get("device"))
                                found = (
                                    imp_product_names.get(
                                        str(device_name).strip().lower(), []
                                    )
                                    if device_name else []
                                )
                                if len(found) == 1:
                                    product_id = found[0]
                                else:
                                    imp_errors.append({
                                        "Excel row": row_no,
                                        "Errors": "Missing or unclear product_id",
                                    })
                                    continue
                            elif imp_key(product_id) not in imp_product_ids:
                                imp_errors.append({
                                    "Excel row": row_no,
                                    "Errors": f"product_id {product_id} does not exist",
                                })
                                continue

                            date_value = clean_value(row.get("price_date"))
                            if date_value is not None:
                                try:
                                    date_value = pd.to_datetime(
                                        date_value
                                    ).date().isoformat()
                                except Exception:
                                    imp_errors.append({
                                        "Excel row": row_no,
                                        "Errors": "price_date is not a valid date",
                                    })
                                    continue

                            record = {"product_id": product_id, "unit_price": price}
                            if supplier_id is not None:
                                record["supplier_id"] = supplier_id
                            if date_value is not None:
                                record["price_date"] = date_value
                            for optional in ("currency", "quote_reference", "notes"):
                                optional_value = clean_value(row.get(optional))
                                if optional_value is not None:
                                    record[optional] = optional_value

                            dup_key = (
                                imp_key(product_id),
                                imp_key(supplier_id),
                                round(price, 6),
                                date_value,
                            )
                            if dup_key in existing_prices or dup_key in seen_prices:
                                skipped_duplicates += 1
                                continue

                            seen_prices.add(dup_key)
                            device_records.append((row_no, record))

                        elif item_type == "accessory":
                            accessory_id = clean_value(row.get("accessory_id"))

                            if accessory_id is None:
                                imp_errors.append({
                                    "Excel row": row_no,
                                    "Errors": "Missing accessory_id",
                                })
                                continue
                            if imp_key(accessory_id) not in imp_acc_ids:
                                imp_errors.append({
                                    "Excel row": row_no,
                                    "Errors": f"accessory_id {accessory_id} does not exist",
                                })
                                continue

                            payload = {"unit_price": price}
                            acc_currency = clean_value(row.get("currency"))
                            if acc_currency is not None:
                                payload["currency"] = (
                                    str(acc_currency).strip().upper()
                                )
                            if supplier_id is not None:
                                payload["supplier_id"] = supplier_id
                            accessory_updates.append(
                                (row_no, accessory_id, payload)
                            )

                        else:
                            imp_errors.append({
                                "Excel row": row_no,
                                "Errors": "item_type must be Device or Accessory",
                            })

                    s1, s2, s3, s4 = st.columns(4)
                    s1.metric("Device prices to add", len(device_records))
                    s2.metric("Accessory prices to update", len(accessory_updates))
                    s3.metric("Skipped (no price)", skipped_empty)
                    s4.metric("Skipped (already saved)", skipped_duplicates)

                    if imp_errors:
                        st.write(f"Rows with errors: {len(imp_errors)}")
                        imp_error_df = pd.DataFrame(imp_errors)
                        st.dataframe(
                            imp_error_df, use_container_width=True, hide_index=True
                        )

                    imp_partial = st.checkbox(
                        "Import valid rows even if some rows contain errors",
                        value=False,
                        key="cp_partial",
                    )

                    imp_can_import = bool(
                        device_records or accessory_updates
                    ) and (not imp_errors or imp_partial)

                    if st.button(
                        "Import prices",
                        type="primary",
                        disabled=not imp_can_import,
                        key="cp_import_btn",
                    ):
                        total_jobs = len(device_records) + len(accessory_updates)
                        progress = st.progress(0)
                        done = 0
                        inserted = 0
                        updated = 0
                        db_errors = []

                        for row_no, record in device_records:
                            try:
                                supabase.table("product_prices").insert(
                                    record
                                ).execute()
                                inserted += 1
                            except Exception as exc:
                                db_errors.append({
                                    "Excel row": row_no, "Error": str(exc)
                                })
                            done += 1
                            progress.progress(done / total_jobs)

                        for row_no, accessory_id, payload in accessory_updates:
                            try:
                                supabase.table("product_accessories").update(
                                    payload
                                ).eq("id", accessory_id).execute()
                                updated += 1
                            except Exception as exc:
                                db_errors.append({
                                    "Excel row": row_no, "Error": str(exc)
                                })
                            done += 1
                            progress.progress(done / total_jobs)

                        st.success(
                            f"Finished. Device prices added: {inserted}; "
                            f"Accessory prices updated: {updated}; "
                            f"Database errors: {len(db_errors)}"
                        )

                        if db_errors:
                            st.dataframe(
                                pd.DataFrame(db_errors),
                                use_container_width=True,
                                hide_index=True,
                            )

            except Exception as exc:
                st.error(f"Import error: {exc}")
# import combined files once.
    
# updated to import one time in one file for products and accessories

# add button to shown products with accessories on system:

    with st.expander("View products with their accessories"):
        vw_products = pd.DataFrame(fetch_rows("products"))
        vw_acc = pd.DataFrame(fetch_rows("product_accessories"))
        vw_prices = pd.DataFrame(fetch_rows("product_prices"))

        def vw_key(v):
            if v is None:
                return None
            try:
                if pd.isna(v):
                    return None
            except (TypeError, ValueError):
                pass
            if isinstance(v, float) and v.is_integer():
                return str(int(v))
            return str(v)

        vw_sup = {
            str(k): v
            for k, v in get_choices("suppliers", "supplier_name").items()
        }

        if vw_products.empty:
            st.info("No products yet.")
        else:
            for col in ["id", "product_name", "category", "manufacturer",
                        "model", "unit"]:
                if col not in vw_products.columns:
                    vw_products[col] = None
            for col in ["product_id", "accessory_name", "accessory_type",
                        "manufacturer", "model", "quantity", "unit",
                        "unit_price", "supplier_id"]:
                if col not in vw_acc.columns:
                    vw_acc[col] = None
            for col in ["product_id", "supplier_id", "unit_price", "currency",
                        "price_date", "created_at"]:
                if col not in vw_prices.columns:
                    vw_prices[col] = None

            w1, w2 = st.columns(2)
            vw_basis = w1.radio(
                "Device price",
                ["Latest price", "Lowest price"],
                horizontal=True,
                key="vw_basis",
            )
            vw_only_acc = w2.checkbox(
                "Only products that have accessories",
                value=True,
                key="vw_only_acc",
            )

            w3, w4 = st.columns(2)
            vw_selected = w3.multiselect(
                "Products (leave empty for all)",
                sorted(vw_products["product_name"].dropna().astype(str).unique()),
                key="vw_selected",
            )
            vw_types = w4.multiselect(
                "Accessory type (leave empty for all)",
                sorted(vw_acc["accessory_type"].dropna().astype(str).unique()),
                key="vw_types",
            )

            vw_products["_k"] = vw_products["id"].map(vw_key)
            vw_acc["_k"] = vw_acc["product_id"].map(vw_key)
            vw_prices["_k"] = vw_prices["product_id"].map(vw_key)

            if vw_types:
                vw_acc = vw_acc[
                    vw_acc["accessory_type"].astype(str).isin(vw_types)
                ]

            vw_prices["unit_price"] = pd.to_numeric(
                vw_prices["unit_price"], errors="coerce"
            )
            vw_prices = vw_prices.dropna(subset=["unit_price"]).copy()
            vw_prices["supplier_name"] = vw_prices["supplier_id"].map(
                lambda v: vw_sup.get(vw_key(v))
            )

            if vw_prices.empty:
                vw_best = pd.DataFrame(columns=[
                    "_k", "device_price", "currency", "device_supplier",
                ])
            else:
                if vw_basis == "Latest price":
                    vw_best = vw_prices.sort_values(
                        ["price_date", "created_at"], na_position="first"
                    ).drop_duplicates("_k", keep="last")
                else:
                    vw_best = vw_prices.sort_values(
                        "unit_price"
                    ).drop_duplicates("_k", keep="first")
                vw_best = vw_best[
                    ["_k", "unit_price", "currency", "supplier_name"]
                ].rename(columns={
                    "unit_price": "device_price",
                    "supplier_name": "device_supplier",
                })

            vw_devices = vw_products.merge(vw_best, on="_k", how="left")

            if vw_selected:
                vw_devices = vw_devices[
                    vw_devices["product_name"].astype(str).isin(vw_selected)
                ]
            if vw_only_acc:
                vw_devices = vw_devices[
                    vw_devices["_k"].isin(set(vw_acc["_k"].dropna()))
                ]

            if vw_devices.empty:
                st.info("No products match these options.")
            else:
                vw_detail = []
                vw_summary = []

                for _, dev in vw_devices.sort_values(
                    "product_name", na_position="last"
                ).iterrows():
                    dev_price = dev.get("device_price")
                    dev_has_price = pd.notna(dev_price)
                    dev_currency = dev.get("currency")

                    vw_detail.append({
                        "device": dev["product_name"],
                        "item_type": "Device",
                        "item_name": dev["product_name"],
                        "accessory_type": None,
                        "manufacturer": dev["manufacturer"],
                        "model": dev["model"],
                        "quantity": 1,
                        "unit": dev["unit"],
                        "unit_price": dev_price if dev_has_price else None,
                        "currency": dev_currency,
                        "line_total": dev_price if dev_has_price else None,
                        "price_supplier": dev.get("device_supplier"),
                    })

                    acc_total = 0.0
                    acc_count = 0
                    unpriced = 0 if dev_has_price else 1

                    for _, acc in vw_acc[vw_acc["_k"] == dev["_k"]].iterrows():
                        qty = pd.to_numeric(acc["quantity"], errors="coerce")
                        price = pd.to_numeric(acc["unit_price"], errors="coerce")
                        qty_used = 1 if pd.isna(qty) else qty

                        if pd.notna(price):
                            line_total = qty_used * price
                            acc_total += line_total
                        else:
                            line_total = None
                            unpriced += 1

                        acc_count += 1
                        vw_detail.append({
                            "device": dev["product_name"],
                            "item_type": "Accessory",
                            "item_name": acc["accessory_name"],
                            "accessory_type": acc["accessory_type"],
                            "manufacturer": acc["manufacturer"],
                            "model": acc["model"],
                            "quantity": None if pd.isna(qty) else qty,
                            "unit": acc["unit"],
                            "unit_price": None if pd.isna(price) else price,
                            "currency": dev_currency,
                            "line_total": line_total,
                            "price_supplier": vw_sup.get(
                                vw_key(acc["supplier_id"])
                            ),
                        })

                    vw_summary.append({
                        "device": dev["product_name"],
                        "category": dev["category"],
                        "device_price": dev_price if dev_has_price else None,
                        "accessories_count": acc_count,
                        "accessories_total": acc_total if acc_count else None,
                        "complete_total": (
                            (dev_price if dev_has_price else 0) + acc_total
                        ),
                        "currency": dev_currency,
                        "items_without_price": unpriced,
                    })

                vw_detail_df = pd.DataFrame(vw_detail)
                vw_summary_df = pd.DataFrame(vw_summary)

                m1, m2, m3 = st.columns(3)
                m1.metric("Devices", len(vw_summary_df))
                m2.metric(
                    "Accessory lines", len(vw_detail_df) - len(vw_summary_df)
                )
                m3.metric(
                    "Items without price",
                    int(vw_summary_df["items_without_price"].sum()),
                )

                st.write("Summary per device")
                st.dataframe(
                    vw_summary_df, use_container_width=True, hide_index=True
                )

                st.write("Devices with their accessories")
                st.dataframe(
                    vw_detail_df, use_container_width=True, hide_index=True
                )

# add button to shown products with accessories on system.
    
    st.divider()

    # price_rows = fetch_rows("product_prices")
    # if not price_rows:
    #     st.info("No prices saved yet.")
    #     st.stop()

    # price_df = pd.DataFrame(price_rows)
    # price_df.insert(
    #     1, "product_name", price_df["product_id"].map(price_products)
    # )
    # price_df.insert(
    #     3, "supplier_name", price_df["supplier_id"].map(price_suppliers)
    # )

    # view = st.radio(
    #     "View",
    #     ["Latest price per product and supplier", "All price records"],
    #     horizontal=True,
    # )

    # if view == "Latest price per product and supplier":
    #     price_df = (
    #         price_df.sort_values(
    #             ["price_date", "created_at"], na_position="first"
    #         )
    #         .drop_duplicates(
    #             subset=["product_id", "supplier_id"], keep="last"
    #         )
    #         .sort_values(["product_name", "supplier_name"])
    #     )

    # p1, p2 = st.columns(2)

    # chosen_products = p1.multiselect(
    #     "Product", sorted(price_df["product_name"].dropna().unique())
    # )
    # chosen_suppliers = p2.multiselect(
    #     "Supplier", sorted(price_df["supplier_name"].dropna().unique())
    # )
    # price_keyword = st.text_input("Search in prices")

    # if chosen_products:
    #     price_df = price_df[price_df["product_name"].isin(chosen_products)]
    # if chosen_suppliers:
    #     price_df = price_df[price_df["supplier_name"].isin(chosen_suppliers)]
    # if price_keyword:
    #     price_mask = (
    #         price_df.astype(str)
    #         .apply(
    #             lambda c: c.str.contains(
    #                 price_keyword, case=False, na=False, regex=False
    #             )
    #         )
    #         .any(axis=1)
    #     )
    #     price_df = price_df[price_mask]

    # q1, q2 = st.columns(2)
    # q1.metric("Price records", len(price_df))
    # q2.metric("Products with a price", price_df["product_id"].nunique())

    # st.dataframe(price_df, use_container_width=True, hide_index=True)

    # st.download_button(
    #     "Download prices (CSV)",
    #     price_df.to_csv(index=False).encode("utf-8-sig"),
    #     file_name="product_prices.csv",
    #     mime="text/csv",
    #     key="csv_product_prices",
    # )

#checkbox if not expander:

    if not st.checkbox(
        "Show price records table", value=False, key="show_price_records"
    ):
        st.stop()
#checkbox if not expander

    products_df = pd.DataFrame(fetch_rows("products"))
    prices_df = pd.DataFrame(fetch_rows("product_prices"))

    if products_df.empty:
        st.info("No products yet.")
        st.stop()

    for col in ["id", "product_name", "category", "manufacturer",
                "model", "unit"]:
        if col not in products_df.columns:
            products_df[col] = None

    for col in ["product_id", "supplier_id", "unit_price", "currency",
                "price_date", "quote_reference", "notes", "created_at"]:
        if col not in prices_df.columns:
            prices_df[col] = None

    def price_key(series):
        return series.map(
            lambda v: None if pd.isna(v)
            else str(int(v)) if isinstance(v, float) and v.is_integer()
            else str(v)
        )

    v1, v2 = st.columns(2)
    view = v1.radio(
        "View",
        ["Latest price per product and supplier", "All price records"],
        horizontal=True,
    )
    include_no_price = v2.checkbox("Also show products without a price")

    prices_df["supplier_name"] = prices_df["supplier_id"].map(price_suppliers)
    prices_df["_pkey"] = price_key(prices_df["product_id"])
    prices_df["_skey"] = price_key(prices_df["supplier_id"])

    if view == "Latest price per product and supplier" and not prices_df.empty:
        prices_df = (
            prices_df.sort_values(
                ["price_date", "created_at"], na_position="first"
            )
            .drop_duplicates(subset=["_pkey", "_skey"], keep="last")
        )

    products_small = products_df[[
        "id", "product_name", "category", "manufacturer", "model", "unit",
    ]].copy()
    products_small["_pkey"] = price_key(products_small["id"])
    products_small = products_small.drop(columns=["id"])

    price_view = prices_df.merge(
        products_small,
        on="_pkey",
        how="right" if include_no_price else "inner",
    )

    show_columns = [
        "product_name", "category", "manufacturer", "model", "unit",
        "supplier_name", "unit_price", "currency", "price_date",
        "quote_reference", "notes",
    ]
    price_view = price_view[show_columns].sort_values(
        ["product_name", "supplier_name"], na_position="last"
    )

    if price_view.empty:
        st.info("No prices saved yet.")
        st.stop()

    p1, p2, p3 = st.columns(3)
    chosen_products = p1.multiselect(
        "Product", sorted(price_view["product_name"].dropna().unique())
    )
    chosen_categories = p2.multiselect(
        "Category", sorted(price_view["category"].dropna().astype(str).unique())
    )
    chosen_suppliers = p3.multiselect(
        "Supplier", sorted(price_view["supplier_name"].dropna().unique())
    )
    price_keyword = st.text_input("Search in prices")

    if chosen_products:
        price_view = price_view[price_view["product_name"].isin(chosen_products)]
    if chosen_categories:
        price_view = price_view[
            price_view["category"].astype(str).isin(chosen_categories)
        ]
    if chosen_suppliers:
        price_view = price_view[
            price_view["supplier_name"].isin(chosen_suppliers)
        ]
    if price_keyword:
        price_mask = (
            price_view.astype(str)
            .apply(
                lambda c: c.str.contains(
                    price_keyword, case=False, na=False, regex=False
                )
            )
            .any(axis=1)
        )
        price_view = price_view[price_mask]

    q1, q2 = st.columns(2)
    q1.metric("Rows", len(price_view))
    q2.metric("Products with a price", price_view.loc[
        price_view["unit_price"].notna(), "product_name"
    ].nunique())

    st.dataframe(price_view, use_container_width=True, hide_index=True)

    st.download_button(
        "Download prices (CSV)",
        price_view.to_csv(index=False).encode("utf-8-sig"),
        file_name="product_prices.csv",
        mime="text/csv",
        key="csv_product_prices",
    )

#Add product price.


elif page == "Inventory":
    add_basic_form(
        "inventory",
        [
            ("product_id", "product"), ("serial_number", "text"),
            ("quantity", "number"), ("stock_classification", "stock"),
            ("project_id", "project"), ("supplier_id", "supplier"),
            ("po_number", "text"), ("receipt_date", "date"),
            ("location", "text"), ("condition", "text"),
            ("notes", "textarea"),
        ],
        "Add Inventory Record",
    )
    display_table("inventory")


elif page == "Purchase Orders":
    add_basic_form(
        "purchase_orders",
        [
            ("po_number", "text"), ("project_id", "project"),
            ("supplier_id", "supplier"), ("order_date", "date"),
            ("currency", "text"), ("status", "text"),
            ("payment_status", "text"), ("shipping_status", "text"),
            ("notes", "textarea"),
        ],
        "Add Purchase Order",
    )
    display_table("purchase_orders")


elif page == "Purchase Order Items":
    add_basic_form(
        "purchase_order_items",
        [
            ("purchase_order_id", "po"), ("product_id", "product"),
            ("quantity", "number"), ("unit_price", "number"),
            ("notes", "textarea"),
        ],
        "Add Purchase Order Item",
    )
    display_table("purchase_order_items")


# -----------------------------------------------------
# Feature: SEARCH any section
# -----------------------------------------------------

elif page == "Search":
    st.subheader("Search")

    mode = st.radio(
        "What do you want to find?",
        [
            "Search any section",
            "Suppliers by category (with products & prices)",
        ],
        horizontal=True,
    )

    # -----------------------------------------------------
    # MODE 1: SEARCH ANY SECTION
    # -----------------------------------------------------
    if mode == "Search any section":
        section = st.selectbox(
            "Section", list(TABLES.keys()), key="search_section"
        )
        table = TABLES[section]["table"]

        rows = fetch_rows(table)
        if not rows:
            st.info("No records found in this section.")
            st.stop()

        df = pd.DataFrame(rows)

        # Show names next to ID columns so results are readable
        search_lookups = {
            "supplier_id": ("suppliers", "supplier_name", "supplier_name"),
            "project_id": ("projects", "project_name", "project_name"),
            "product_id": ("products", "product_name", "product_name"),
            "purchase_order_id": ("purchase_orders", "po_number", "po_number"),
        }

        for fk, (ref_table, label_col, new_name) in search_lookups.items():
            if fk in df.columns and new_name not in df.columns:
                try:
                    names = get_choices(ref_table, label_col)
                except Exception:
                    continue
                df.insert(
                    df.columns.get_loc(fk) + 1,
                    new_name,
                    df[fk].map(names),
                )

        keyword = st.text_input(
            "Search in all columns", key="search_keyword"
        )

        filter_cols = st.multiselect(
            "Filter by column (optional)",
            list(df.columns),
            key="search_filter_cols",
        )

        filtered = df.copy()

        for col in filter_cols:
            values = filtered[col].dropna().astype(str).unique()
            if len(values) <= 50:
                chosen = st.multiselect(
                    f"{col}", sorted(values), key=f"search_f_{section}_{col}"
                )
                if chosen:
                    filtered = filtered[
                        filtered[col].astype(str).isin(chosen)
                    ]
            else:
                text = st.text_input(
                    f"{col} contains", key=f"search_t_{section}_{col}"
                )
                if text:
                    filtered = filtered[
                        filtered[col].astype(str).str.contains(
                            text, case=False, na=False, regex=False
                        )
                    ]

        if keyword:
            mask = (
                filtered.astype(str)
                .apply(
                    lambda c: c.str.contains(
                        keyword, case=False, na=False, regex=False
                    )
                )
                .any(axis=1)
            )
            filtered = filtered[mask]

        c1, c2 = st.columns(2)
        c1.metric("Records found", len(filtered))
        c2.metric("Total records in section", len(df))

        st.dataframe(filtered, use_container_width=True, hide_index=True)

        st.download_button(
            "Download results (CSV)",
            filtered.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"search_{table}.csv",
            mime="text/csv",
            key="search_csv_any",
        )

    # -----------------------------------------------------
    # MODE 2: SUPPLIERS BY CATEGORY WITH PRODUCTS & PRICES
    # -----------------------------------------------------
    else:
        def search_ensure(df, cols):
            for c in cols:
                if c not in df.columns:
                    df[c] = None
            return df

        suppliers = pd.DataFrame(fetch_rows("suppliers"))
        if suppliers.empty:
            st.info("No suppliers yet.")
            st.stop()

        suppliers = search_ensure(suppliers, [
            "id", "supplier_name", "category", "products", "contact_person",
            "contact_title", "email", "phone", "payment_terms",
        ])

        categories = sorted({
            str(c).strip() for c in suppliers["category"].dropna()
            if str(c).strip()
        })

        col1, col2 = st.columns(2)
        category = col1.selectbox(
            "Category", ["All categories"] + categories, key="search_category"
        )
        keyword2 = col2.text_input(
            "Keyword (supplier, product, contact...)", key="search_keyword2"
        )

        filtered_sup = suppliers.copy()
        if category != "All categories":
            filtered_sup = filtered_sup[
                filtered_sup["category"].astype(str).str.strip().str.lower()
                == category.lower()
            ]

        if filtered_sup.empty:
            st.info("No suppliers found in this category.")
            st.stop()

        products = search_ensure(pd.DataFrame(fetch_rows("products")), [
            "id", "product_name", "manufacturer", "model", "unit",
            "supplier_id",
        ])
        pos = search_ensure(pd.DataFrame(fetch_rows("purchase_orders")), [
            "id", "po_number", "supplier_id", "order_date", "currency",
            "status",
        ])
        items = search_ensure(
            pd.DataFrame(fetch_rows("purchase_order_items")),
            ["purchase_order_id", "product_id", "quantity", "unit_price"],
        )

        sup_small = filtered_sup[[
            "id", "supplier_name", "category", "products", "contact_person",
            "contact_title", "email", "phone", "payment_terms",
        ]].rename(columns={"id": "sid", "products": "products_listed"})

        po_small = pos[[
            "id", "po_number", "supplier_id", "order_date", "currency",
            "status",
        ]].rename(columns={"id": "po_id", "supplier_id": "sid"})

        item_small = items[[
            "purchase_order_id", "product_id", "quantity", "unit_price",
        ]]

        prod_small = products[[
            "id", "product_name", "manufacturer", "model", "unit",
            "supplier_id",
        ]].rename(columns={"id": "product_id", "supplier_id": "prod_sid"})

        # A) Products that have prices from purchase orders
        priced = item_small.merge(
            po_small, left_on="purchase_order_id", right_on="po_id",
            how="inner",
        )
        priced = priced.merge(
            prod_small.drop(columns=["prod_sid"]), on="product_id", how="left"
        )
        priced = priced.merge(sup_small, on="sid", how="inner")

        # B) Products linked to the supplier but never ordered (no price yet)
        linked = prod_small.rename(columns={"prod_sid": "sid"}).merge(
            sup_small, on="sid", how="inner"
        )
        priced_pairs = set(zip(priced["sid"], priced["product_id"]))
        linked = linked[[
            (s, p) not in priced_pairs
            for s, p in zip(linked["sid"], linked["product_id"])
        ]]

        # C) Suppliers with no products or prices at all
        covered = set(priced["sid"]) | set(linked["sid"])
        no_data = sup_small[~sup_small["sid"].isin(covered)]

        result = pd.concat([priced, linked, no_data], ignore_index=True)

        result["line_total"] = (
            pd.to_numeric(result["quantity"], errors="coerce")
            * pd.to_numeric(result["unit_price"], errors="coerce")
        )

        display_columns = [
            "supplier_name", "category", "contact_person", "contact_title",
            "email", "phone", "payment_terms", "products_listed",
            "product_name", "manufacturer", "model", "unit",
            "po_number", "order_date", "currency", "quantity",
            "unit_price", "line_total", "status",
        ]
        result = result[display_columns].sort_values(
            ["supplier_name", "product_name"], na_position="last"
        )

        if keyword2:
            mask = (
                result.astype(str)
                .apply(
                    lambda c: c.str.contains(
                        keyword2, case=False, na=False, regex=False
                    )
                )
                .any(axis=1)
            )
            result = result[mask]

        m1, m2, m3 = st.columns(3)
        m1.metric("Suppliers", result["supplier_name"].nunique())
        m2.metric("Products", result["product_name"].nunique())
        m3.metric("Price records", int(result["unit_price"].notna().sum()))

        st.write("Suppliers, products and prices")
        st.dataframe(result, use_container_width=True, hide_index=True)

        st.download_button(
            "Download results (CSV)",
            result.to_csv(index=False).encode("utf-8-sig"),
            file_name="supplier_search.csv",
            mime="text/csv",
            key="search_csv_suppliers",
        )

        priced_rows = result.dropna(subset=["unit_price"]).copy()
        if not priced_rows.empty:
            priced_rows["unit_price"] = pd.to_numeric(
                priced_rows["unit_price"], errors="coerce"
            )
            priced_rows = priced_rows.sort_values("order_date")

            summary = (
                priced_rows
                .groupby(
                    ["product_name", "supplier_name", "currency"],
                    dropna=False,
                )["unit_price"]
                .agg(lowest="min", highest="max", latest="last", orders="count")
                .reset_index()
                .sort_values(["product_name", "lowest"])
            )

            st.write("Price summary (lowest, highest and latest per product)")
            st.dataframe(summary, use_container_width=True, hide_index=True)




# =========================================================
# EXCEL IMPORT
# =========================================================

elif page == "Excel Import":
    st.subheader("Excel / CSV Import")

    # ---------- Template download ----------
    with st.expander("Download Excel template for a section", expanded=True):
        template_label = st.selectbox(
            "Section", list(TABLES.keys()), key="template_section"
        )
        template_with_id = st.checkbox(
            "Include 'id' column (only for updating existing records)",
            value=False,
            key="template_with_id",
        )

        try:
            template_bytes = build_template(template_label, template_with_id)
            st.download_button(
                f"Download {template_label} template",
                template_bytes,
                file_name=(
                    f"{TABLES[template_label]['table']}_template.xlsx"
                ),
                mime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                key="template_download",
            )
            st.caption(
                "Fill the 'Template' sheet only. Yellow headers are "
                "required. The 'Instructions' and '... ref' sheets show "
                "the rules and the IDs to use for linked fields. "
                "Then upload the file below."
            )
        except Exception as exc:
            st.error(f"Could not build template: {exc}")

    st.divider()

    uploaded = st.file_uploader(
        "Upload Excel or CSV file",
        type=["xlsx", "xls", "csv"],
    )

    if uploaded:
        try:
            content = uploaded.getvalue()

            if uploaded.name.lower().endswith(".csv"):
                source_df = pd.read_csv(io.BytesIO(content))
            else:
                excel = pd.ExcelFile(io.BytesIO(content))
                sheet = st.selectbox("Worksheet", excel.sheet_names)
                source_df = pd.read_excel(
                    io.BytesIO(content), sheet_name=sheet
                )

            source_df.columns = [
                str(column).strip() for column in source_df.columns
            ]

            st.write("Source preview")
            st.dataframe(source_df.head(10), use_container_width=True)

            target_label = st.selectbox("Destination table", list(TABLES.keys()))
            config = TABLES[target_label]
            table = config["table"]

            st.write("Map the Excel columns to database fields.")

            mapping = {}
            options = ["-- Ignore --"] + config["columns"]
            columns = st.columns(2)

            for i, source_column in enumerate(source_df.columns):
                suggestion = suggest_mapping(source_column, config)
                with columns[i % 2]:
                    mapping[source_column] = st.selectbox(
                        f"{source_column}",
                        options,
                        index=options.index(suggestion)
                        if suggestion in options else 0,
                        key=f"mapping_{table}_{i}",
                    )

            targets = [
                value for value in mapping.values()
                if value != "-- Ignore --"
            ]

            duplicates = sorted({
                value for value in targets
                if targets.count(value) > 1
            })

            if duplicates:
                st.error(
                    "More than one source column maps to: "
                    + ", ".join(duplicates)
                )
                st.stop()

            mapped = pd.DataFrame()

            for source_column, target_column in mapping.items():
                if target_column != "-- Ignore --":
                    mapped[target_column] = source_df[source_column]

            mapped = mapped.astype(object).where(pd.notna(mapped), None)
            mapped = mapped.dropna(how="all").reset_index(drop=True)
    # added for empty rows "unit price in importing       
            if table == "product_prices" and "unit_price" in mapped.columns:
                mapped = mapped[mapped["unit_price"].notna()].reset_index(drop=True)
    # added for empty rows "unit price in importing  
            if mapped.empty:
                st.warning("No mapped data found.")
                st.stop()

            mode = st.radio(
                "Import mode",
                ["Insert", "Update existing", "Upsert"],
                horizontal=True,
            )

 # to ignore empty cells and update other completed           
            if (
                table == "product_accessories"
                and mode == "Update existing"
                and "unit_price" in mapped.columns
            ):
                mapped = mapped[mapped["unit_price"].notna()].reset_index(drop=True)
# to ignore empty cells and update other completed  
            match_key = None

            if mode == "Update existing":
                available_keys = [
                    col for col in mapped.columns
                    if col in config["columns"]
                ]
                if not available_keys:
                    st.error("Map a unique key such as ID or PO number.")
                    st.stop()

                default_key = config["key"]
                match_key = st.selectbox(
                    "Match existing records by",
                    available_keys,
                    index=available_keys.index(default_key)
                    if default_key in available_keys else 0,
                )

            elif mode == "Upsert":
                match_key = config["key"]
                if match_key not in mapped.columns:
                    st.error(
                        f"Upsert requires the '{match_key}' column. "
                        "Map that column or use Insert."
                    )
                    st.stop()

            records = []
            validation_errors = []

# Add product price:
            price_product_ids = {}
            price_supplier_ids = {}
            if table == "product_prices":
                for pid, pname in get_choices(
                    "products", "product_name"
                ).items():
                    price_product_ids.setdefault(
                        str(pname).strip().lower(), []
                    ).append(pid)
                for sid, sname in get_choices(
                    "suppliers", "supplier_name"
                ).items():
                    price_supplier_ids[str(sname).strip().lower()] = sid
# Add product price.
            
#added with product accessories:
            product_ids_by_name = {}
            if table == "product_accessories":
                for pid, pname in get_choices(
                    "products", "product_name"
                ).items():
                    product_ids_by_name.setdefault(
                        str(pname).strip().lower(), []
                    ).append(pid)
#added with product accessories.
            
            
            for row_index, (_, row) in enumerate(mapped.iterrows(), start=2):
                record = clean_record(row.to_dict())
                row_errors = []

 #Add product price:               
                if table == "product_prices":
                    price_product_name = record.pop("product_name", None)
                    price_supplier_name = record.pop("supplier_name", None)

                    if record.get("product_id") is None and price_product_name:
                        matches = price_product_ids.get(
                            str(price_product_name).strip().lower(), []
                        )
                        if len(matches) == 1:
                            record["product_id"] = matches[0]
                        elif not matches:
                            row_errors.append(
                                f"Product not found: {price_product_name}"
                            )
                        else:
                            row_errors.append(
                                f"Several products named {price_product_name}; "
                                "use product_id instead"
                            )

                    if record.get("supplier_id") is None and price_supplier_name:
                        found_supplier = price_supplier_ids.get(
                            str(price_supplier_name).strip().lower()
                        )
                        if found_supplier is None:
                            row_errors.append(
                                f"Supplier not found: {price_supplier_name}"
                            )
                        else:
                            record["supplier_id"] = found_supplier

#Add product price.
                
                
#added with product accessories:
                if table == "product_accessories":
                    product_name_value = record.pop("product_name", None)
                    if record.get("product_id") is None and product_name_value:
                        matches = product_ids_by_name.get(
                            str(product_name_value).strip().lower(), []
                        )
                        if len(matches) == 1:
                            record["product_id"] = matches[0]
                        elif not matches:
                            row_errors.append(
                                f"Product not found: {product_name_value}"
                            )
                        else:
                            row_errors.append(
                                f"Several products named {product_name_value}; "
                                "use product_id instead"
                            )
                    if record.get("product_id") is None and not product_name_value:
                        row_errors.append(
                            "Missing product_name or product_id"
                        )

                for field in config["required"]:
                    if record.get(field) is None:
                        row_errors.append(f"Missing {field}")

                if mode != "Insert" and record.get(match_key) is None:
                    row_errors.append(f"Missing matching key: {match_key}")

                if table == "inventory":
                    stock_type = record.get("stock_classification")
                    if stock_type is not None and stock_type not in STOCK_TYPES:
                        row_errors.append(
                            "Invalid stock_classification. Use: "
                            + ", ".join(STOCK_TYPES)
                        )
#added with product accessories.

                
                for number_field in ("quantity", "unit_price"):
                    if record.get(number_field) is not None:
                        try:
                            number = Decimal(str(record[number_field]))
                            if number < 0:
                                row_errors.append(
                                    f"{number_field} cannot be negative"
                                )
                        except (InvalidOperation, ValueError):
                            row_errors.append(
                                f"{number_field} must be numeric"
                            )

                if row_errors:
                    validation_errors.append({
                        "Excel row": row_index,
                        "Errors": "; ".join(row_errors),
                    })
                else:
                    records.append((row_index, record))

            # Duplicate matching keys in the same file
            if mode != "Insert" and match_key:
                seen_keys = set()
                filtered_records = []

                for row_index, record in records:
                    key_value = str(record[match_key]).strip().lower()

                    if key_value in seen_keys:
                        validation_errors.append({
                            "Excel row": row_index,
                            "Errors": f"Duplicate {match_key}: {record[match_key]}",
                        })
                    else:
                        seen_keys.add(key_value)
                        filtered_records.append((row_index, record))

                records = filtered_records

            st.subheader("Mapped data preview")
            st.dataframe(
                pd.DataFrame([record for _, record in records]),
                use_container_width=True,
            )

            st.write(f"Valid rows: {len(records)}")
            st.write(f"Rows with validation errors: {len(validation_errors)}")

            if validation_errors:
                error_df = pd.DataFrame(validation_errors)
                st.dataframe(error_df, use_container_width=True)

                st.download_button(
                    "Download validation errors",
                    error_df.to_csv(index=False).encode("utf-8-sig"),
                    file_name="validation_errors.csv",
                    mime="text/csv",
                )

            partial_import = st.checkbox(
                "Import valid rows even if some rows contain errors",
                value=False,
            )

            can_import = bool(records) and (
                not validation_errors or partial_import
            )

            if st.button("Start import", type="primary", disabled=not can_import):
                inserted = 0
                updated = 0
                skipped = 0
                database_errors = []

                progress = st.progress(0)

                for position, (row_index, record) in enumerate(records):
                    payload = dict(record)
                    payload.pop("created_at", None)

                    try:
                        if mode == "Insert":
                            # Let Supabase generate the ID unless one was provided.
                            supabase.table(table).insert(payload).execute()
                            inserted += 1

                        elif mode == "Update existing":
                            key_value = payload.pop(match_key)

                            existing = (
                                supabase.table(table)
                                .select("id")
                                .eq(match_key, key_value)
                                .limit(1)
                                .execute()
                                .data
                                or []
                            )

                            if not existing:
                                skipped += 1
                                database_errors.append({
                                    "Excel row": row_index,
                                    "Error": f"No record found for {match_key}={key_value}",
                                })
                            else:
                                supabase.table(table).update(payload).eq(
                                    match_key, key_value
                                ).execute()
                                updated += 1

                        else:
                            supabase.table(table).upsert(
                                payload,
                                on_conflict=match_key,
                            ).execute()
                            updated += 1

                    except Exception as exc:
                        database_errors.append({
                            "Excel row": row_index,
                            "Error": str(exc),
                        })

                    progress.progress((position + 1) / len(records))

                st.success(
                    f"Finished. Inserted: {inserted}; "
                    f"Updated/Upserted: {updated}; Skipped: {skipped}; "
                    f"Database errors: {len(database_errors)}"
                )

                if database_errors:
                    db_error_df = pd.DataFrame(database_errors)
                    st.dataframe(db_error_df, use_container_width=True)

                    st.download_button(
                        "Download database errors",
                        db_error_df.to_csv(index=False).encode("utf-8-sig"),
                        file_name="database_errors.csv",
                        mime="text/csv",
                    )

        except Exception as exc:
            st.error(f"Import error: {exc}")


# =========================================================
# DOCUMENT MANAGEMENT
# =========================================================

elif page == "Documents":
    st.subheader("Documents")

    entity_types = {
        "Supplier": ("suppliers", "supplier_name", "supplier_id"),
        "Project": ("projects", "project_name", "project_id"),
        "Product": ("products", "product_name", "product_id"),
        "Inventory": ("inventory", "serial_number", "inventory_id"),
        "Purchase Order": ("purchase_orders", "po_number", "purchase_order_id"),
    }

    entity_type = st.selectbox("Link document to", list(entity_types.keys()))
    entity_table, label_field, foreign_key = entity_types[entity_type]

    try:
        entity_rows = fetch_rows(entity_table)
        entity_options = {
            row["id"]: str(row.get(label_field) or f"ID {row['id']}")
            for row in entity_rows
        }

        if not entity_options:
            st.warning("Create a record before uploading a document.")
        else:
            selected_id = st.selectbox(
                "Select record",
                list(entity_options.keys()),
                format_func=lambda x: entity_options[x],
            )

            documents = st.file_uploader(
                "Select documents",
                accept_multiple_files=True,
            )

            if st.button(
                "Upload documents",
                disabled=not documents,
            ):
                for document in documents:
                    try:
                        safe_name = re.sub(
                            r"[^A-Za-z0-9._-]+", "_", document.name
                        )
                        path = (
                            f"{entity_type.lower().replace(' ', '_')}/"
                            f"{selected_id}/"
                            f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_"
                            f"{safe_name}"
                        )

                        supabase.storage.from_(BUCKET).upload(
                            path,
                            document.getvalue(),
                            {
                                "content-type": document.type
                                or "application/octet-stream"
                            },
                        )

                        metadata = {
                            "file_name": document.name,
                            "storage_path": path,
                            "file_type": document.type,
                            foreign_key: selected_id,
                        }

                        supabase.table("attachments").insert(
                            metadata
                        ).execute()

                        st.success(f"Uploaded {document.name}")

                    except Exception as exc:
                        st.error(f"Upload failed for {document.name}: {exc}")

        st.divider()
        st.subheader("Existing documents")

        attachments = (
            supabase.table("attachments")
            .select("*")
            .order("id", desc=True)
            .limit(1000)
            .execute()
            .data
            or []
        )

        if attachments:
            st.dataframe(
                pd.DataFrame(attachments),
                use_container_width=True,
                hide_index=True,
            )

            attachment_by_id = {
                row["id"]: row for row in attachments
            }

            selected_attachment_id = st.selectbox(
                "Choose document",
                list(attachment_by_id.keys()),
                format_func=lambda x: (
                    f"{attachment_by_id[x].get('file_name', 'Document')} "
                    f"(ID {x})"
                ),
            )

            if st.button("Create download link"):
                selected = attachment_by_id[selected_attachment_id]
                signed = supabase.storage.from_(BUCKET).create_signed_url(
                    selected["storage_path"], 3600
                )

                url = signed.get("signedURL") or signed.get("signedUrl")

                if url:
                    st.markdown(f"[Open / Download document]({url})")
                else:
                    st.error("Could not generate download link.")
        else:
            st.info("No documents uploaded yet.")

    except Exception as exc:
        st.error(f"Document section error: {exc}")
